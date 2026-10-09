from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("CreatorTwin_AI_Creator_Marketplace_Pitch.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "CreatorTwin | AI Creator Marketplace"
prs.core_properties.subject = "10-slide product and differentiation pitch"
prs.core_properties.author = "CreatorTwin"
prs.core_properties.keywords = "AI creators, marketplace, campaign briefs, provenance, licensing"

SW, SH = 13.333, 7.5
NAVY = "0B1720"
NAVY2 = "132630"
INK = "14252D"
BODY = "42545C"
MUTED = "73838A"
CREAM = "F7F7F2"
WHITE = "FFFFFF"
TEAL = "0F766E"
MINT = "5EEAD4"
CORAL = "FF6B5E"
LAV = "B8B2FF"
LINE = "DFE6E3"
PALE_TEAL = "E7F5F2"
PALE_CORAL = "FFF0ED"
PALE_LAV = "F0EEFF"
FONT = "Aptos"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def rect(slide, x, y, w, h, fill, radius=False, line=None, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        try:
            shape.adjustments[0] = 0.12
        except Exception:
            pass
    return shape


def circle(slide, x, y, d, fill, line=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5, dash=None):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    return shape


def txt(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP, font=FONT, margin=0, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, part in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        p.space_before = Pt(0)
        run = p.add_run()
        run.text = part
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
    return box


def rich_link_text(slide, parts, x, y, w, h, size=8, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    for label, url in parts:
        run = p.add_run()
        run.text = label
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(TEAL if url else color)
        if url:
            run.hyperlink.address = url
    return box


def pill(slide, label, x, y, w, fill=PALE_TEAL, color=TEAL, h=0.30, size=9):
    rect(slide, x, y, w, h, fill, radius=True)
    txt(slide, label, x + 0.08, y + 0.01, w - 0.16, h - 0.02, size=size, color=color,
        bold=True, valign=MSO_ANCHOR.MIDDLE)


def base(bg=CREAM):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def header(slide, section, title, subtitle=None, dark=False):
    main = WHITE if dark else INK
    sub = "B7C8CD" if dark else MUTED
    accent = MINT if dark else TEAL
    txt(slide, section.upper(), 0.65, 0.38, 8.7, 0.24, size=9, color=accent, bold=True)
    txt(slide, title, 0.65, 0.72, 11.9, 0.56, size=28, color=main, bold=True)
    if subtitle:
        txt(slide, subtitle, 0.68, 1.34, 11.8, 0.42, size=12, color=sub)


def footer(slide, n, dark=False, source_parts=None):
    c = "82969D" if dark else "89979A"
    line(slide, 0.65, 7.08, 12.68, 7.08, color="28404A" if dark else LINE, width=0.7)
    txt(slide, "CREATORTWIN  /  AI CREATOR MARKETPLACE", 0.66, 7.16, 4.7, 0.18, size=7.5, color=c, bold=True)
    if source_parts:
        rich_link_text(slide, source_parts, 5.20, 7.14, 6.65, 0.24, size=7, color=c)
    txt(slide, f"{n:02d} / 10", 11.92, 7.14, 0.72, 0.2, size=7.5, color=c, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, fill=WHITE, stroke=LINE):
    return rect(slide, x, y, w, h, fill, radius=True, line=stroke)


def bullet(slide, label, x, y, w, color=BODY, dot=TEAL, size=13, h=0.52):
    circle(slide, x, y + 0.10, 0.09, dot)
    txt(slide, label, x + 0.22, y, w - 0.22, h, size=size, color=color)


# 1 — Cover
s = base(NAVY)
circle(s, 9.8, -0.4, 4.5, TEAL, transparency=66)
circle(s, 10.9, 3.7, 3.4, CORAL, transparency=72)
circle(s, 8.8, 4.6, 2.2, LAV, transparency=79)
pill(s, "AI-NATIVE CREATOR MARKETPLACE", 0.72, 0.67, 2.72, fill=NAVY2, color=MINT, h=0.35, size=8.5)
txt(s, "AI creators are here.\nHiring systems aren’t ready.", 0.72, 1.42, 7.8, 1.55, size=34, color=WHITE, bold=True)
txt(s, "CreatorTwin turns AI-creator discovery into an evidence-led\nworkflow—from campaign brief to approved delivery.",
    0.76, 3.28, 7.0, 0.86, size=17, color="D3E0E3")
line(s, 0.78, 4.55, 6.7, 4.55, color="39515A", width=1)
txt(s, "DISCOVER", 0.78, 4.79, 1.35, 0.25, size=9, color=MINT, bold=True)
txt(s, "BRIEF", 2.28, 4.79, 1.1, 0.25, size=9, color=MINT, bold=True)
txt(s, "MATCH", 3.66, 4.79, 1.1, 0.25, size=9, color=MINT, bold=True)
txt(s, "DELIVER", 5.05, 4.79, 1.3, 0.25, size=9, color=MINT, bold=True)
for px in [1.78, 3.31, 4.70]:
    txt(s, "→", px, 4.76, 0.35, 0.28, size=14, color=CORAL, bold=True, align=PP_ALIGN.CENTER)

# floating product evidence card
card(s, 8.48, 1.39, 3.95, 4.36, NAVY2, stroke="34505B")
circle(s, 8.82, 1.75, 0.59, TEAL)
txt(s, "MJ", 8.82, 1.88, 0.59, 0.20, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, "Manindra Jonnakuti", 9.58, 1.77, 2.25, 0.30, size=15, color=WHITE, bold=True)
txt(s, "AI filmmaker · Guntur, India", 9.59, 2.12, 2.4, 0.25, size=9, color="B8C7D0")
line(s, 8.82, 2.63, 12.09, 2.63, color="34505B", width=0.8)
txt(s, "PORTFOLIO EVIDENCE", 8.84, 2.84, 2.45, 0.22, size=8, color=MINT, bold=True)
txt(s, "AI cinematic short film", 8.84, 3.19, 3.08, 0.32, size=16, color=WHITE, bold=True)
pill(s, "VEO 3.1 FAST", 8.84, 3.67, 1.23, fill="21474B", color=MINT, size=8)
pill(s, "FLOW AI", 10.19, 3.67, 0.82, fill="303A56", color=LAV, size=8)
txt(s, "Workflow", 8.84, 4.20, 1.2, 0.22, size=9, color="B8C7D0", bold=True)
txt(s, "Concept  →  AI production  →  final film", 8.84, 4.49, 3.06, 0.3, size=10, color=WHITE)
pill(s, "RIGHTS UNVERIFIED", 8.84, 5.07, 1.48, fill="513732", color="FFB2A8", size=8)
txt(s, "A visible gap becomes a visible signal.", 0.75, 6.38, 5.2, 0.3, size=12, color="AFC2C7", italic=True)
footer(s, 1, dark=True)


# 2 — Problem
s = base()
header(s, "01 / The problem", "A beautiful output is not enough to hire with confidence.",
       "AI-native work introduces questions a thumbnail alone cannot answer.")
pain_cards = [
    ("01", "Discovery", "Who can make this with the right model, medium and style?", TEAL, PALE_TEAL),
    ("02", "Evaluation", "What tools and workflow sit behind the portfolio piece?", "6758C8", PALE_LAV),
    ("03", "Commercial use", "Can the brand actually use the delivered work—and where?", CORAL, PALE_CORAL),
    ("04", "Collaboration", "Where do versions, revision notes and approval live?", "AB6A12", "FFF5E5"),
]
for i, (num, title, body, color, pale) in enumerate(pain_cards):
    x = 0.68 + (i % 2) * 6.06
    y = 2.02 + (i // 2) * 1.58
    card(s, x, y, 5.72, 1.32)
    circle(s, x + 0.24, y + 0.24, 0.46, pale)
    txt(s, num, x + 0.24, y + 0.34, 0.46, 0.17, size=9, color=color, bold=True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.87, y + 0.21, 4.4, 0.29, size=16, color=INK, bold=True)
    txt(s, body, x + 0.87, y + 0.61, 4.45, 0.47, size=12, color=BODY)
txt(s, "TODAY’S HANDOFF", 0.72, 5.52, 1.55, 0.22, size=8.5, color=MUTED, bold=True)
steps = [("Portfolio link", 0.72), ("DM / email", 3.14), ("Loose brief", 5.56), ("Unclear rights", 7.98), ("File + feedback", 10.4)]
for label, x in steps:
    rect(s, x, 5.91, 1.82, 0.56, PALE_CORAL, radius=True)
    txt(s, label, x + 0.1, 6.07, 1.62, 0.19, size=9, color="9B514A", bold=True, align=PP_ALIGN.CENTER)
for x in [2.58, 5.0, 7.42, 9.84]:
    txt(s, "…", x, 6.04, 0.4, 0.25, size=18, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
footer(s, 2)


# 3 — Category comparison
s = base()
header(s, "02 / The landscape", "Adjacent products solve parts of the journey.",
       "CreatorTwin is designed around the missing connective layer for AI campaign work.")
cols = [
    ("FREELANCE MARKETPLACES", "Find talent · scope projects · manage contracts", "Strong at general hiring and transactions.", "AI production evidence and rights fit are not the organizing unit.", TEAL, PALE_TEAL),
    ("PORTFOLIO NETWORKS", "Showcase · browse · follow creative work", "Strong at visual discovery and inspiration.", "A campaign brief and its delivery/review state live elsewhere.", "6758C8", PALE_LAV),
    ("AI CREATION TOOLS", "Generate · edit · iterate with models", "Strong at making content with AI tools.", "They are creation environments, not creator marketplaces.", CORAL, PALE_CORAL),
]
for i, (title, subtitle, strength, gap, accent, pale) in enumerate(cols):
    x = 0.68 + i * 4.22
    card(s, x, 2.02, 3.92, 3.58)
    rect(s, x, 2.02, 3.92, 0.14, accent, radius=True)
    txt(s, title, x + 0.22, 2.34, 3.45, 0.27, size=10, color=accent, bold=True)
    txt(s, subtitle, x + 0.22, 2.75, 3.45, 0.55, size=14, color=INK, bold=True)
    txt(s, "STRENGTH", x + 0.22, 3.54, 1.0, 0.18, size=8, color=MUTED, bold=True)
    txt(s, strength, x + 0.22, 3.80, 3.45, 0.48, size=11, color=BODY)
    txt(s, "OPEN GAP FOR THIS USE CASE", x + 0.22, 4.50, 2.72, 0.18, size=8, color=MUTED, bold=True)
    rect(s, x + 0.17, 4.82, 3.58, 0.57, pale, radius=True)
    txt(s, gap, x + 0.32, 4.90, 3.28, 0.42, size=10, color=INK)
rect(s, 0.70, 5.88, 11.96, 0.63, NAVY, radius=True)
txt(s, "CREATORTWIN", 0.94, 6.07, 1.6, 0.2, size=9, color=MINT, bold=True)
txt(s, "Evidence-led creator discovery  +  structured brief  +  rights-aware match  +  delivery workflow",
    2.65, 6.03, 9.4, 0.29, size=12, color=WHITE, bold=True)
footer(s, 3, source_parts=[("Upwork hiring overview  ·  ", "https://support.upwork.com/hc/en-us/articles/211063398-How-to-start-hiring-on-Upwork"),
                           ("Behance overview  ·  ", "https://help.behance.net/hc/en-us/articles/204483894-What-is-Behance"),
                           ("Adobe Firefly", "https://www.adobe.com/products/firefly.html")])


# 4 — Product
s = base()
header(s, "03 / The solution", "One workflow from “find a creator” to “approve the work.”",
       "CreatorTwin turns AI production metadata into practical hiring and collaboration decisions.")
flow = [
    ("01", "DISCOVER", "Filter by specialty, tools, skills, format and portfolio evidence.", TEAL, PALE_TEAL),
    ("02", "BRIEF", "Turn a rough idea into a structured campaign brief.", "6758C8", PALE_LAV),
    ("03", "MATCH", "See what aligns, what is missing and whether rights fit.", CORAL, PALE_CORAL),
    ("04", "DELIVER", "Track delivery links, notes, revisions, approval and review.", "AB6A12", "FFF5E5"),
]
for i, (num, title, body, accent, pale) in enumerate(flow):
    x = 0.70 + i * 3.16
    card(s, x, 2.40, 2.76, 2.75)
    circle(s, x + 0.23, 2.68, 0.55, pale)
    txt(s, num, x + 0.23, 2.84, 0.55, 0.18, size=10, color=accent, bold=True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.23, 3.47, 2.30, 0.28, size=12, color=accent, bold=True)
    txt(s, body, x + 0.23, 3.91, 2.28, 0.89, size=12, color=BODY)
    if i < 3:
        circle(s, x + 2.87, 3.57, 0.30, NAVY)
        txt(s, "→", x + 2.87, 3.62, 0.30, 0.16, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(s, 0.70, 5.52, 11.92, 0.90, PALE_TEAL, radius=True)
txt(s, "THE PRODUCT PROMISE", 0.96, 5.78, 1.93, 0.2, size=8, color=TEAL, bold=True)
txt(s, "Make the creative fit legible before the contract—and the work reviewable after it.",
    3.00, 5.73, 8.9, 0.33, size=16, color=INK, bold=True)
footer(s, 4)


# 5 — User journeys
s = base()
header(s, "04 / Who it serves", "Three roles. One shared source of truth.",
       "Designed for the brand buying the work, the creator making it, and the reviewer validating trust signals.")
roles = [
    ("BRAND / AGENCY", "“Can this creator deliver this brief—and can we use the result?”", ["Build a brief", "Compare fit reasons", "Review delivery + revisions"], TEAL, PALE_TEAL),
    ("AI CREATOR", "“Can I prove how I work and get a fair, clear brief?”", ["Show tools + workflow", "Set rights and formats", "Submit versioned work"], "6758C8", PALE_LAV),
    ("REVIEWER", "“Is this claim backed by evidence?”", ["Inspect submitted evidence", "Approve or request changes", "Record a decision"], CORAL, PALE_CORAL),
]
for i, (role, quote, bullets, accent, pale) in enumerate(roles):
    x = 0.69 + i * 4.23
    card(s, x, 2.08, 3.93, 3.64)
    circle(s, x + 0.25, 2.34, 0.48, pale)
    txt(s, ["B", "C", "R"][i], x + 0.25, 2.48, 0.48, 0.17, size=12, color=accent, bold=True, align=PP_ALIGN.CENTER)
    txt(s, role, x + 0.89, 2.43, 2.7, 0.23, size=9, color=accent, bold=True)
    txt(s, quote, x + 0.25, 3.03, 3.36, 0.66, size=14, color=INK, bold=True)
    for j, item in enumerate(bullets):
        bullet(s, item, x + 0.28, 4.00 + j * 0.48, 3.28, size=10.5, h=0.38)
rect(s, 0.72, 6.00, 11.90, 0.51, NAVY, radius=True)
txt(s, "BRIEF OWNER", 1.02, 6.16, 1.3, 0.16, size=8, color=MINT, bold=True)
txt(s, "Brand / agency", 2.31, 6.12, 1.43, 0.2, size=10, color=WHITE)
txt(s, "→", 3.86, 6.11, 0.28, 0.2, size=11, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, "WORK OWNER", 4.33, 6.16, 1.33, 0.16, size=8, color=MINT, bold=True)
txt(s, "Creator", 5.70, 6.12, 1.1, 0.2, size=10, color=WHITE)
txt(s, "→", 7.06, 6.11, 0.28, 0.2, size=11, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, "TRUST OWNER", 7.51, 6.16, 1.36, 0.16, size=8, color=MINT, bold=True)
txt(s, "Reviewer", 8.92, 6.12, 1.4, 0.2, size=10, color=WHITE)
txt(s, "→ shared evidence", 10.20, 6.12, 2.0, 0.2, size=9, color="C8D6DA")
footer(s, 5)


# 6 — AI-native proof layer
s = base()
header(s, "05 / The differentiator", "Make the process behind the pixels visible.",
       "AI portfolios carry the production context brands need—not just the final image.")
card(s, 0.70, 2.00, 7.18, 4.55)
pill(s, "PUBLIC PORTFOLIO PREVIEW", 0.98, 2.28, 2.05, size=8)
txt(s, "AI cinematic short film", 0.98, 2.78, 4.8, 0.40, size=20, color=INK, bold=True)
txt(s, "Creator-attributed · source-linked · not independently verified", 0.99, 3.24, 5.9, 0.25, size=10, color=MUTED)
line(s, 0.98, 3.73, 7.51, 3.73)
txt(s, "TOOLS / MODEL", 0.99, 3.98, 1.53, 0.19, size=8, color=MUTED, bold=True)
pill(s, "Google Flow AI", 0.99, 4.31, 1.34, fill=PALE_LAV, color="594BAB", size=8)
pill(s, "Google Veo 3.1 Fast", 2.43, 4.31, 1.74, fill=PALE_TEAL, color=TEAL, size=8)
pill(s, "After Effects", 4.27, 4.31, 1.19, fill="F2F3F2", color=BODY, size=8)
txt(s, "WORKFLOW", 0.99, 4.91, 1.2, 0.18, size=8, color=MUTED, bold=True)
txt(s, "Concept  →  AI-assisted production  →  cinematic presentation", 0.99, 5.20, 6.2, 0.32, size=12, color=INK, bold=True)
pill(s, "AI USE: CREATOR-DESCRIBED", 0.99, 5.82, 2.21, fill=PALE_TEAL, color=TEAL, size=8)
pill(s, "RIGHTS: UNVERIFIED", 3.33, 5.82, 1.78, fill=PALE_CORAL, color="A34D43", size=8)
txt(s, "Source: Manindra Jonnakuti’s public Behance project. These sample records are attributed previews, not rights-cleared CreatorTwin inventory.",
    0.99, 6.25, 6.3, 0.20, size=7.5, color=MUTED)

txt(s, "EVIDENCE LABELS", 8.30, 2.12, 2.2, 0.24, size=9, color=TEAL, bold=True)
evidence = [
    ("CREATOR-DESCRIBED", "Claim is attributed to the public creator source.", PALE_TEAL, TEAL),
    ("REVIEWER-VERIFIED", "Reserved for a claim approved through review.", PALE_LAV, "594BAB"),
    ("NOT DISCLOSED", "Missing tools or models stay visibly missing.", PALE_CORAL, "A34D43"),
    ("RIGHTS UNVERIFIED", "Public visibility is not a usage license.", "FFF5E5", "AB6A12"),
]
for i, (tag, desc, pale, color) in enumerate(evidence):
    yy = 2.56 + i * 0.91
    pill(s, tag, 8.30, yy, 1.80 if i != 1 else 1.84, fill=pale, color=color, size=7.7)
    txt(s, desc, 10.27, yy + 0.01, 2.26, 0.51, size=10, color=BODY)
rect(s, 8.30, 6.24, 4.24, 0.30, NAVY, radius=True)
txt(s, "Portfolio item = media + provenance + workflow + rights", 8.45, 6.31, 3.95, 0.15, size=8.5, color=WHITE, bold=True)
footer(s, 6, source_parts=[("Behance project ↗", "https://www.behance.net/gallery/252637499/AI-Filmmaker-VFX-Artist-Cinematic-Storytelling-")])


# 7 — Brief builder
s = base()
header(s, "06 / Brief definition", "Turn a fuzzy idea into a brief a creator can act on.",
       "The AI assistant structures the request; the brand remains the decision-maker.")
card(s, 0.70, 2.00, 4.24, 4.61, NAVY, stroke=NAVY)
pill(s, "BRAND IDEA", 1.00, 2.30, 1.00, fill=NAVY2, color=MINT, size=8)
txt(s, "“We’re launching a new\nenergy drink. Need something\ncinematic for social.”", 1.00, 2.88, 3.50, 1.37, size=22, color=WHITE, bold=True)
line(s, 1.00, 4.60, 4.63, 4.60, color="34505B")
txt(s, "AI ASSISTANT", 1.00, 4.86, 1.40, 0.19, size=8, color=MINT, bold=True)
txt(s, "Asks for the missing choices: audience, aspect ratio, runtime and usage rights.",
    1.00, 5.20, 3.55, 0.68, size=12, color="D1DFE2")
pill(s, "GEMINI-POWERED WHEN CONFIGURED", 1.00, 6.10, 2.70, fill="24343B", color="B8C7D0", size=7)
card(s, 5.25, 2.00, 7.38, 4.61)
txt(s, "STRUCTURED CAMPAIGN BRIEF", 5.57, 2.32, 3.60, 0.22, size=9, color=TEAL, bold=True)
brief_fields = [
    ("OBJECTIVE", "Launch awareness"), ("CONTENT", "AI short-form film"),
    ("STYLE", "Cinematic · high energy"), ("AUDIENCE", "College students"),
    ("FORMAT", "9:16 · 30 seconds"), ("TOOLS / SKILLS", "Veo · editing · storyboarding"),
    ("USAGE", "Paid social · commercial license required"), ("DELIVERY", "Instagram · due date + budget"),
]
for idx, (lab, val) in enumerate(brief_fields):
    col, row = idx % 2, idx // 2
    x, y = 5.57 + col * 3.47, 2.84 + row * 0.82
    rect(s, x, y, 3.16, 0.66, "F0F7F5" if row % 2 == 0 else "F6F6F3", radius=True)
    txt(s, lab, x + 0.13, y + 0.11, 2.9, 0.15, size=7.5, color=MUTED, bold=True)
    txt(s, val, x + 0.13, y + 0.32, 2.90, 0.22, size=10, color=INK, bold=True)
txt(s, "Clear inputs improve matching and prevent late surprises about format, process or rights.",
    5.58, 6.23, 6.6, 0.2, size=9, color=BODY, italic=True)
footer(s, 7)


# 8 — Matching
s = base()
header(s, "07 / Discovery & filtering", "A fit score should explain itself.",
       "CreatorTwin compares only the requirements in the brief and shows matched and missing evidence.")
card(s, 0.70, 2.08, 3.12, 4.43, NAVY, stroke=NAVY)
txt(s, "BRIEF FILTERS", 0.99, 2.39, 1.85, 0.22, size=9, color=MINT, bold=True)
for i, (label, value) in enumerate([
    ("SPECIALTY", "AI filmmaking"), ("FORMAT", "9:16 video"),
    ("TOOLS", "Veo / Flow"), ("LICENSE", "Paid commercial use"),
]):
    yy = 2.91 + i * 0.69
    txt(s, label, 0.99, yy, 1.16, 0.16, size=7.5, color="90A8AF", bold=True)
    txt(s, value, 0.99, yy + 0.22, 2.43, 0.22, size=11, color=WHITE, bold=True)
    if i < 3:
        line(s, 0.99, yy + 0.54, 3.48, yy + 0.54, color="28404A", width=0.8)
pill(s, "RIGHTS ARE A REAL REQUIREMENT", 0.99, 5.89, 2.46, fill="513732", color="FFB2A8", size=7)

txt(s, "ILLUSTRATIVE MATCH EXPLANATION", 4.20, 2.16, 3.45, 0.20, size=8, color=MUTED, bold=True)
matches = [
    ("Creator A", "Skills + tool + format align", "License evidence is missing → verify before hire", "84", TEAL, PALE_TEAL),
    ("Creator B", "Style + format align", "Required tool not listed · license scope differs", "62", "AB6A12", "FFF5E5"),
    ("Creator C", "Strong visual specialty", "No commercial-use claim → not a rights fit", "—", CORAL, PALE_CORAL),
]
for i, (name, good, warn, score, accent, pale) in enumerate(matches):
    yy = 2.50 + i * 1.22
    card(s, 4.18, yy, 8.45, 1.02)
    circle(s, 4.42, yy + 0.28, 0.45, pale)
    txt(s, score, 4.42, yy + 0.40, 0.45, 0.16, size=9, color=accent, bold=True, align=PP_ALIGN.CENTER)
    txt(s, name, 5.03, yy + 0.15, 1.65, 0.22, size=13, color=INK, bold=True)
    txt(s, good, 6.65, yy + 0.14, 5.47, 0.22, size=10, color=TEAL)
    txt(s, warn, 6.65, yy + 0.51, 5.53, 0.25, size=9.5, color=accent, bold=True)
rect(s, 4.19, 6.27, 8.43, 0.27, PALE_CORAL, radius=True)
txt(s, "A generic “commercial use available” claim cannot stand in for a required license scope.",
    4.37, 6.32, 8.07, 0.17, size=9, color="914E46", bold=True)
txt(s, "Illustrative scores only—not live CreatorTwin rankings.", 4.22, 5.99, 4.2, 0.17, size=8, color=MUTED, italic=True)
footer(s, 8)


# 9 — Lifecycle + verification
s = base()
header(s, "08 / From match to outcome", "A “delivered” status should mean work is ready to review.",
       "The engagement record keeps the artifact, revision context and decision together.")
stages = [
    ("01", "Request", "brief + fit"), ("02", "Accept", "scope + date"),
    ("03", "Produce", "creator work"), ("04", "Deliver", "URL + notes"),
    ("05", "Revise", "round + notes"), ("06", "Approve", "review + rating"),
]
for i, (num, name, desc) in enumerate(stages):
    x = 0.70 + i * 2.10
    accent = TEAL if i in [0, 3, 5] else (CORAL if i == 4 else "59717A")
    pale = PALE_TEAL if i in [0, 3, 5] else (PALE_CORAL if i == 4 else "EDF1EF")
    card(s, x, 2.51, 1.74, 1.56)
    circle(s, x + 0.18, 2.72, 0.42, pale)
    txt(s, num, x + 0.18, 2.84, 0.42, 0.15, size=8.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    txt(s, name, x + 0.18, 3.27, 1.35, 0.25, size=12, color=INK, bold=True)
    txt(s, desc, x + 0.18, 3.62, 1.40, 0.20, size=9, color=BODY)
    if i < 5:
        txt(s, "→", x + 1.78, 3.14, 0.28, 0.21, size=12, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
rect(s, 0.70, 4.53, 7.70, 1.68, NAVY, radius=True)
txt(s, "DELIVERY RECORD", 0.98, 4.82, 1.65, 0.2, size=8.5, color=MINT, bold=True)
pill(s, "Version 02", 0.99, 5.25, 1.02, fill=NAVY2, color=WHITE, size=8)
txt(s, "delivery link", 2.20, 5.26, 1.2, 0.18, size=9, color="D8E3E5")
txt(s, "+", 3.49, 5.21, 0.27, 0.24, size=14, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, "creator notes", 3.84, 5.26, 1.25, 0.18, size=9, color="D8E3E5")
txt(s, "+", 5.25, 5.21, 0.27, 0.24, size=14, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, "revision history", 5.61, 5.26, 1.48, 0.18, size=9, color="D8E3E5")
rect(s, 8.68, 4.53, 3.94, 1.68, PALE_LAV, radius=True)
txt(s, "SEPARATE TRUST LANE", 8.97, 4.82, 2.4, 0.2, size=8.5, color="594BAB", bold=True)
txt(s, "Creator submits evidence", 8.97, 5.18, 2.7, 0.21, size=11, color=INK, bold=True)
txt(s, "Reviewer approves or requests changes → visible signal", 8.97, 5.52, 3.18, 0.39, size=9.5, color=BODY)
txt(s, "Prototype flow: URL delivery is supported; binary file storage is not yet wired up.",
    0.72, 6.54, 8.6, 0.18, size=8, color=MUTED, italic=True)
footer(s, 9)


# 10 — Why now / close
s = base(NAVY)
circle(s, 10.2, -0.5, 4.0, TEAL, transparency=73)
circle(s, 11.0, 4.9, 2.9, CORAL, transparency=78)
header(s, "09 / Why CreatorTwin", "Not another place to post AI art.",
       "A workflow layer that makes AI creative work easier to evaluate, commission and review.", dark=True)
reasons = [
    ("AI-native evidence", "Models, tools, workflow and formats are part of the portfolio data."),
    ("Brief-to-fit logic", "A structured brief drives field-level matching with readable reasons."),
    ("Rights-aware", "License scope is a separate requirement—not a vague badge."),
    ("Accountable delivery", "Delivery links, notes, revisions and approval stay connected."),
]
for i, (title, body) in enumerate(reasons):
    col, row = i % 2, i // 2
    x, y = 0.72 + col * 4.06, 2.15 + row * 1.42
    circle(s, x, y + 0.05, 0.42, TEAL if i in [0, 2] else CORAL)
    txt(s, ["01", "02", "03", "04"][i], x, y + 0.18, 0.42, 0.13, size=7.7, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.59, y + 0.01, 3.15, 0.27, size=15, color=WHITE, bold=True)
    txt(s, body, x + 0.59, y + 0.40, 3.22, 0.58, size=10.5, color="B7C8CD")
card(s, 9.00, 2.11, 3.62, 3.22, NAVY2, stroke="34505B")
txt(s, "THE WEDGE", 9.31, 2.42, 1.4, 0.19, size=8.5, color=MINT, bold=True)
txt(s, "Evidence\n× Brief\n× Rights\n× Delivery", 9.31, 2.85, 2.90, 1.76, size=24, color=WHITE, bold=True)
pill(s, "PROTOTYPE · NOT MARKETPLACE TRACTION", 9.31, 4.82, 2.95, fill="24343B", color="CAD7DA", size=7)
line(s, 0.74, 5.52, 12.56, 5.52, color="34505B")
txt(s, "NEXT PROOF", 0.76, 5.81, 1.16, 0.18, size=8, color=MINT, bold=True)
txt(s, "Pilot with opt-in creators and real briefs. Measure shortlist time, rights clarity and revision cycles.",
    2.00, 5.76, 9.9, 0.34, size=13, color=WHITE, bold=True)
txt(s, "The defensible idea is the connected evidence model—not a claim that no other platform has any of these features.",
    0.76, 6.49, 10.9, 0.21, size=9, color="AFC2C7", italic=True)
footer(s, 10, dark=True)


prs.save(OUT)
print(f"Saved {len(prs.slides)} slides to {OUT}")
